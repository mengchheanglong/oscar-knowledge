"""Bounded, non-executing local extractors. Invoked in a killable child process."""

import csv
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from bs4 import BeautifulSoup

from app.config import Settings

PARSER_VERSION = "segments-v1"
SUPPORTED = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".html",
    ".htm",
    ".srt",
    ".vtt",
    ".csv",
    ".xlsx",
    ".bookpart",
}


class SourceError(ValueError):
    pass


@dataclass
class ExtractedSegment:
    text: str
    heading_path: str | None = None
    page_start: int | None = None
    page_end: int | None = None
    start_ms: int | None = None
    end_ms: int | None = None
    metadata: dict = field(default_factory=dict)


def validate_zip(path: Path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(x.file_size for x in entries) > 50_000_000:
            raise SourceError("ARCHIVE_LIMIT")
        for entry in entries:
            if (
                entry.filename.startswith(("/", "\\"))
                or ".." in Path(entry.filename).parts
                or ":" in entry.filename
                or entry.file_size / max(entry.compress_size, 1) > 200
            ):
                raise SourceError("UNSAFE_ARCHIVE")


def timestamp(value: str) -> int:
    parts = value.replace(",", ".").split(":")
    if len(parts) == 2:
        parts.insert(0, "0")
    h, m, s = parts
    return round((int(h) * 3600 + int(m) * 60 + float(s)) * 1000)


def subtitles(text: str) -> list[ExtractedSegment]:
    result: list[ExtractedSegment] = []
    previous: list[str] = []
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        lines = block.splitlines()
        timing = next((i for i, line in enumerate(lines) if "-->" in line), None)
        if timing is None:
            continue
        match = re.search(r"([\d:.,]+)\s+-->\s+([\d:.,]+)", lines[timing])
        if not match:
            raise SourceError("INVALID_SUBTITLE")
        start, end = map(timestamp, match.groups())
        if start < 0 or end < start:
            raise SourceError("INVALID_SUBTITLE")
        words = BeautifulSoup(" ".join(lines[timing + 1 :]), "html.parser").get_text().split()
        overlap = 0
        for n in range(1, min(len(previous), len(words)) + 1):
            if previous[-n:] == words[:n]:
                overlap = n
        original = words
        words = words[overlap:]
        if words:
            result.append(ExtractedSegment(" ".join(words), start_ms=start, end_ms=end))
        previous = original
    return result


def load(path: Path, settings: Settings) -> list[ExtractedSegment]:
    if path.is_symlink() or not path.is_file():
        raise SourceError("INVALID_FILE")
    if path.stat().st_size > settings.max_file_bytes:
        raise SourceError("FILE_TOO_LARGE")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise SourceError("UNSUPPORTED_FORMAT")
    result: list[ExtractedSegment] = []

    total_chars = 0

    def add(segment: ExtractedSegment):
        nonlocal total_chars
        if segment.text.strip():
            result.append(segment)
            total_chars += len(segment.text)
        if len(result) > 50000 or total_chars > settings.max_extract_chars:
            raise SourceError("EXTRACTION_LIMIT")

    if suffix == ".bookpart":
        import json

        if path.stat().st_size > 3 * 1024 * 1024:
            raise SourceError("BOOKPART_BYTE_LIMIT")
        data = json.loads(path.read_text(encoding="utf-8"))
        if data.get("format") != "oscar-bookpart-v1" or not 1 <= len(data["pages"]) <= 8:
            raise SourceError("BOOKPART_INVALID")
        for item in data["pages"]:
            if item["page"] is not None and (not isinstance(item["page"], int) or item["page"] < 1):
                raise SourceError("BOOKPART_INVALID")
            add(
                ExtractedSegment(
                    item["text"],
                    heading_path=item.get("chapter"),
                    page_start=item["page"],
                    page_end=item["page"],
                    metadata={**data["metadata"], **item.get("metadata", {})},
                )
            )
    elif suffix == ".pdf":
        import pymupdf

        with path.open("rb") as source:
            signature = source.read(1024)
        if not signature.lstrip().startswith(b"%PDF-"):
            raise SourceError("INVALID_PDF")
        with pymupdf.open(path) as pdf:
            if pdf.needs_pass or pdf.page_count > 2000:
                raise SourceError("PDF_LIMIT_OR_ENCRYPTED")
            for i, page in enumerate(pdf):
                add(ExtractedSegment(page.get_text(), page_start=i + 1, page_end=i + 1))
        if not result:
            raise SourceError("OCR_REQUIRED")
    elif suffix in {".docx", ".xlsx"}:
        validate_zip(path)
        if suffix == ".docx":
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            heading = "Document"
            for block in Document(str(path)).iter_inner_content():
                if isinstance(block, Paragraph):
                    if block.style and block.style.name.startswith("Heading"):
                        heading = block.text
                    add(ExtractedSegment(block.text, heading_path=heading))
                elif isinstance(block, Table):
                    add(
                        ExtractedSegment(
                            "\n".join(" | ".join(c.text for c in row.cells) for row in block.rows),
                            heading_path=heading,
                            metadata={"table": True},
                        )
                    )
        else:
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book:
                    headers = None
                    for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if i > 20000 or len(row) > 100:
                            raise SourceError("CATALOG_LIMIT")
                        if headers is None:
                            headers = [str(v or "") for v in row]
                        else:
                            add(
                                ExtractedSegment(
                                    "; ".join(
                                        f"{h}: {v}" for h, v in zip(headers, row) if v is not None
                                    ),
                                    heading_path=f"{sheet.title} row {i}",
                                    metadata={"row": i},
                                )
                            )
            finally:
                book.close()
    else:
        try:
            text = path.read_text(encoding="utf-8-sig", errors="strict")
        except UnicodeError as exc:
            raise SourceError("ENCODING_UTF8_REQUIRED") from exc
        if "\x00" in text:
            raise SourceError("BINARY_CONTENT")
        if suffix in {".html", ".htm"}:
            soup = BeautifulSoup(text, "html.parser")
            for node in soup(["script", "style", "nav", "header", "footer", "aside", "form"]):
                node.decompose()
            main = soup.find("main") or soup.find("article") or soup
            add(ExtractedSegment(main.get_text("\n", strip=True), heading_path="HTML main"))
        elif suffix in {".srt", ".vtt"}:
            for segment in subtitles(text):
                add(segment)
        elif suffix == ".csv":
            csv.field_size_limit(100000)
            with path.open(encoding="utf-8-sig", newline="") as stream:
                for i, row in enumerate(csv.DictReader(stream), 2):
                    if i > 20000 or len(row) > 100:
                        raise SourceError("CATALOG_LIMIT")
                    add(
                        ExtractedSegment(
                            "; ".join(f"{k}: {v}" for k, v in row.items()),
                            heading_path=f"Row {i}",
                            metadata={"row": i},
                        )
                    )
        else:
            heading = "Text"
            for part in re.split(r"\n\s*\n", text):
                lines = part.splitlines()
                if lines and lines[0].startswith("#"):
                    heading = lines[0].lstrip("# ")
                add(ExtractedSegment(part, heading_path=heading))
    if not result:
        raise SourceError("EMPTY_SOURCE")
    return result
