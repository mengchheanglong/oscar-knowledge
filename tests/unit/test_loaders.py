import zipfile

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from app.config import Settings
from app.ingestion.chunking import chunk_segments
from app.ingestion.loaders import ExtractedSegment, SourceError, load


def test_pdf_pages_and_image_only(tmp_path):
    path = tmp_path / "lecture.pdf"
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Page one subscription pricing")
    pdf.new_page().insert_text((72, 72), "Page two competitive position")
    pdf.save(path)
    segments = load(path, Settings())
    assert [s.page_start for s in segments] == [1, 2]
    pdf.close()
    empty = pymupdf.open()
    empty.new_page()
    empty.save(tmp_path / "image.pdf")
    empty.close()
    with pytest.raises(SourceError, match="OCR_REQUIRED"):
        load(tmp_path / "image.pdf", Settings())


def test_docx_preserves_headings_table_order(tmp_path):
    path = tmp_path / "lecture.docx"
    doc = Document()
    doc.add_heading("Pricing", 1)
    doc.add_paragraph("Before table")
    doc.add_table(1, 1).cell(0, 0).text = "Margin table"
    doc.add_paragraph("After table")
    doc.save(path)
    segments = load(path, Settings())
    assert [s.text for s in segments] == ["Pricing", "Before table", "Margin table", "After table"]
    assert all(s.heading_path == "Pricing" for s in segments)


@pytest.mark.parametrize("suffix", [".txt", ".md", ".html", ".srt", ".vtt", ".csv", ".xlsx"])
def test_formats(tmp_path, suffix):
    path = tmp_path / ("lecture" + suffix)
    content = {
        ".txt": "Pricing text",
        ".md": "# Pricing\n\nPricing text",
        ".html": "<nav>REMOVE</nav><main><h1>Pricing</h1><script>REMOVE</script><p>Pricing text</p></main>",
        ".srt": "1\n00:00:01,000 --> 00:00:03,000\nPricing text\n\n2\n00:00:03,000 --> 00:00:05,000\nPricing text continued\n",
        ".vtt": "WEBVTT\n\n00:01.000 --> 00:03.000\nPricing text\n",
        ".csv": "Topic,Text\nPricing,Pricing text\n",
    }
    if suffix == ".xlsx":
        book = Workbook()
        book.active.append(["Topic", "Text"])
        book.active.append(["Pricing", "Pricing text"])
        book.save(path)
    else:
        path.write_text(content[suffix], encoding="utf-8")
    segments = load(path, Settings())
    assert "Pricing" in " ".join(s.text for s in segments)
    assert "REMOVE" not in " ".join(s.text for s in segments)
    if suffix in {".srt", ".vtt"}:
        assert segments[0].start_ms == 1000
        assert segments[0].end_ms == 3000
    if suffix == ".srt":
        assert segments[1].text == "continued"
    if suffix in {".csv", ".xlsx"}:
        assert segments[0].metadata["row"] == 2


def test_malformed_and_limits(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_text("not pdf")
    with pytest.raises(SourceError, match="INVALID_PDF"):
        load(path, Settings())
    path = tmp_path / "bad.txt"
    path.write_bytes(b"\xff")
    with pytest.raises(SourceError, match="ENCODING"):
        load(path, Settings())
    path.write_text("12345")
    with pytest.raises(SourceError, match="FILE_TOO_LARGE"):
        load(path, Settings(max_file_bytes=2))
    with pytest.raises(SourceError, match="EXTRACTION_LIMIT"):
        load(path, Settings(max_extract_chars=2))
    path.write_text("")
    with pytest.raises(SourceError, match="EMPTY_SOURCE"):
        load(path, Settings())
    path = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("../escape.xml", "unsafe")
    with pytest.raises(SourceError, match="UNSAFE_ARCHIVE"):
        load(path, Settings())


def test_chunk_bounds_overlap_and_determinism():
    segment = ExtractedSegment("abcdefghij" * 1000, page_start=2, page_end=2)
    chunks = list(chunk_segments([segment]))
    assert chunks == list(chunk_segments([segment]))
    assert all(len(c.text) <= 2400 and c.page_start == 2 for c in chunks)
    assert chunks[0].text[-360:] == chunks[1].text[:360]
